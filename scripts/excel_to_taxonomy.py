"""Convert F&F labeling Excel to machine-readable taxonomy JSON.

Source : source/F&F_odd key_values_fin.xlsx
Output : st_cut-dev/taxonomy/labeling_taxonomy.json

Excel structure:
  - Sheet `category`: cols [No., cat, sub_cat]   -> 11 main cats x ~21..49 subcats
  - Sheet `detail`  : cols [No., cat, key, value] -> cat -> attribute_key -> value

Detail sheet has:
  - Item categories : Bag, Bottom, Headwear, Hosiery, Inner, Onepiece, Outer, Shoes,
                      Swimwear bottoms, Swimwear inner, Swimwear onepiece
  - Shared bins     : common, model, background, styling

Conventions:
  - Main category names = canonical Title Case (matches category sheet, with Swimwear
    unified across the 3 detail variants).
  - Typos in attribute keys are normalized via TYPO_FIXES.
  - Order is preserved as it appears in the Excel (insertion order).

Usage:
    python st_cut-dev/scripts/excel_to_taxonomy.py
"""
from __future__ import annotations

import json
import sys
from collections import OrderedDict
from datetime import date
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parents[2]
EXCEL_PATH = PROJECT_ROOT / "source" / "F&F_odd key_values_fin.xlsx"
OUT_PATH = PROJECT_ROOT / "st_cut-dev" / "taxonomy" / "labeling_taxonomy.json"

TYPO_FIXES = {
    "sheos detail": "shoes detail",
    "shoes closure tyoe": "shoes closure type",
    "swimwearonepiece upper neckline": "swimwear onepiece upper neckline",
}

SHARED_BINS = {"common", "model", "background", "styling"}

CAT_CANONICAL = {
    "bag": "Bag",
    "bottom": "Bottom",
    "headwear": "Headwear",
    "hosiery": "Hosiery",
    "inner": "Inner",
    "onepiece": "Onepiece",
    "outer": "Outer",
    "shoes": "Shoes",
    "swimwear": "Swimwear",
    "eyewear": "Eyewear",
    "neckwear": "Neckwear",
}

DETAIL_TO_MAIN = {
    "Swimwear bottoms": "Swimwear",
    "Swimwear inner": "Swimwear",
    "Swimwear onepiece": "Swimwear",
}


def normalize_key(key: str) -> str:
    return TYPO_FIXES.get(key, key)


def canonical_cat(cat: str) -> str:
    return CAT_CANONICAL.get(cat.strip().lower(), cat.strip())


def main() -> int:
    if not EXCEL_PATH.exists():
        print(f"[ERROR] Excel not found: {EXCEL_PATH}", file=sys.stderr)
        return 1

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True, read_only=True)

    # ---- Sheet: category -> categories[main_cat].subcategories ----
    categories: "OrderedDict[str, dict]" = OrderedDict()
    ws_cat = wb["category"]
    rows_cat = ws_cat.iter_rows(min_row=2, values_only=True)
    for row in rows_cat:
        if not row or len(row) < 3:
            continue
        _, cat, sub_cat = row[0], row[1], row[2]
        if not cat or not sub_cat:
            continue
        main = canonical_cat(str(cat))
        bucket = categories.setdefault(
            main, {"subcategories": [], "attribute_keys": OrderedDict()}
        )
        sub_norm = str(sub_cat).strip()
        if sub_norm and sub_norm not in bucket["subcategories"]:
            bucket["subcategories"].append(sub_norm)

    # ---- Sheet: detail -> attribute_keys (per cat) + shared ----
    shared: "OrderedDict[str, OrderedDict[str, list]]" = OrderedDict(
        (b, OrderedDict()) for b in ("common", "model", "background", "styling")
    )

    ws_det = wb["detail"]
    rows_det = ws_det.iter_rows(min_row=2, values_only=True)
    for row in rows_det:
        if not row or len(row) < 4:
            continue
        _, cat, key, val = row[0], row[1], row[2], row[3]
        if cat is None or key is None or val is None:
            continue
        cat_s = str(cat).strip()
        key_s = normalize_key(str(key).strip())
        val_s = str(val).strip()
        if not cat_s or not key_s or not val_s:
            continue

        if cat_s in SHARED_BINS:
            target = shared[cat_s].setdefault(key_s, [])
        else:
            main = DETAIL_TO_MAIN.get(cat_s, canonical_cat(cat_s))
            bucket = categories.setdefault(
                main, {"subcategories": [], "attribute_keys": OrderedDict()}
            )
            target = bucket["attribute_keys"].setdefault(key_s, [])
        if val_s not in target:
            target.append(val_s)

    # ---- Stats ----
    n_subcat = sum(len(c["subcategories"]) for c in categories.values())
    all_keys = set()
    all_vals = set()
    for c in categories.values():
        for k, vs in c["attribute_keys"].items():
            all_keys.add(k)
            all_vals.update(vs)
    for bin_name, keys in shared.items():
        for k, vs in keys.items():
            all_keys.add(k)
            all_vals.update(vs)

    payload = {
        "version": date.today().isoformat(),
        "source_excel": EXCEL_PATH.name,
        "stats": {
            "categories": len(categories),
            "subcategories": n_subcat,
            "attribute_keys": len(all_keys),
            "values": len(all_vals),
        },
        "categories": categories,
        "shared": shared,
        "normalization": {"typo_fixes": TYPO_FIXES},
    }

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    with OUT_PATH.open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

    print(f"[OK] wrote {OUT_PATH}")
    print(f"     categories      : {payload['stats']['categories']}")
    print(f"     subcategories   : {payload['stats']['subcategories']}")
    print(f"     attribute_keys  : {payload['stats']['attribute_keys']}")
    print(f"     unique values   : {payload['stats']['values']}")
    print("\n--- per-category attribute_key counts ---")
    for cat, info in categories.items():
        n_keys = len(info["attribute_keys"])
        n_subs = len(info["subcategories"])
        print(f"  {cat:<14} subcats={n_subs:>3}  attr_keys={n_keys:>3}")
    print("\n--- shared bin sizes ---")
    for bin_name, keys in shared.items():
        print(f"  {bin_name:<11} keys={len(keys):>3}")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
